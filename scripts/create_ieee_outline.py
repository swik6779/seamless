from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/home/swik6779/Desktop/seamless/laplacian_blending_ieee_outline.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(header_cells[index], "1F4E79")
        set_cell_border(header_cells[index])
        if widths:
            header_cells[index].width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F2F6FA")
            set_cell_border(cells[index])
            if widths:
                cells[index].width = Inches(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].paragraph_format.space_before = Pt(13)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)
    styles["Heading 3"].font.size = Pt(10.5)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("Comparative Study of Laplacian Pyramid Image Blending")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Serial CPU OpenMP and CUDA Implementations")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(64, 64, 64)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(9)
    intro.add_run("Purpose. ").bold = True
    intro.add_run(
        "This document outlines the main content, experiments, and reporting requirements "
        "for an IEEE-format paper comparing serial CPU, OpenMP, and CUDA implementations "
        "of float32 Laplacian pyramid image blending. The serial implementation serves as "
        "the correctness reference, while OpenMP and CUDA are evaluated for performance, "
        "scalability, and numerical agreement."
    )

    document.add_heading("Proposed Research Questions", level=1)
    for item in [
        "How much speedup does OpenMP provide over the serial CPU implementation?",
        "When does CUDA outperform CPU and OpenMP for the complete blending pipeline?",
        "How much do host-device transfers affect measured CUDA performance?",
        "Which pipeline stages dominate runtime on each backend?",
        "How do Gaussian kernel size and pyramid level count affect runtime?",
        "Do all implementations reproduce the serial reference within float tolerance?",
    ]:
        add_number(document, item)

    document.add_heading("IEEE Paper Structure", level=1)
    add_table(
        document,
        ["Section", "Main points to cover"],
        [
            ("Abstract", "Problem, implementations compared, metrics, and principal findings."),
            ("Index Terms", "Laplacian pyramid, image blending, OpenMP, CUDA, GPU acceleration."),
            ("I. Introduction", "Motivation, problem definition, study objectives, and contributions."),
            ("II. Background", "Gaussian blur, Gaussian/Laplacian pyramids, masks, and multiresolution blending."),
            ("III. Implementation", "Serial CPU, OpenMP, CUDA, data layout, borders, resizing, and memory management."),
            ("IV. Methodology", "Hardware, software, inputs, parameters, timing policy, and repetitions."),
            ("V. Results", "Runtime, stage breakdown, speedups, transfers, throughput, and kernel-size effects."),
            ("VI. Discussion", "Bottlenecks, scaling behavior, GPU crossover point, and tradeoffs."),
            ("VII. Limitations", "Hardware dependence, memory usage, resize policy, and lack of ground-truth blends."),
            ("VIII. Conclusion", "Main findings, practical recommendations, and future work."),
            ("References", "IEEE-numbered citations for Burt-Adelson pyramids, OpenMP, CUDA, and image processing."),
        ],
        widths=[1.3, 5.8],
    )

    document.add_heading("Technical Topics", level=1)
    document.add_heading("Common Algorithm and Data Representation", level=2)
    for item in [
        "Float32 internal buffers; conversion to 8-bit only for final output.",
        "Interleaved row-major channel layout for grayscale, RGB, and RGBA images.",
        "Clamp-to-edge border handling, used consistently by CPU, OpenMP, and CUDA.",
        "Odd-size indexing using ceil(width/2) and ceil(height/2) during downsampling.",
        "Exact target dimensions during bilinear expansion and automatic image alignment.",
        "Image A as the spatial reference when image B or the mask has different dimensions.",
    ]:
        add_bullet(document, item)

    document.add_heading("Backend Implementations", level=2)
    add_table(
        document,
        ["Backend", "Implementation topics"],
        [
            ("Serial CPU", "Reference implementation; separable blur; sequential level construction; correctness baseline."),
            ("OpenMP", "Pixel-level parallelism in blur, downsample, resize, expand, subtract, blend, and collapse; OMP_NUM_THREADS control."),
            ("CUDA", "Device-resident level buffers; kernels for pyramid stages; CUDA events; H2D/kernel/D2H timing."),
        ],
        widths=[1.25, 5.85],
    )

    document.add_heading("Experimental Methodology", level=1)
    document.add_heading("Benchmark Inputs and Parameters", level=2)
    for item in [
        "Use identical in-memory input arrays for every backend in a comparison.",
        "Use a fixed benchmark image size for the main experiment; use additional sizes only if needed for scalability analysis.",
        "Sweep Gaussian kernel sizes 3, 5, 7, 9, and 11.",
        "Sweep pyramid level counts such as 1 through 6, subject to the minimum-side stopping rule.",
        "Evaluate grayscale and RGB inputs if time permits; keep channel count fixed within each comparison.",
        "Use deterministic synthetic images plus a representative apple-orange style image pair and mask.",
    ]:
        add_bullet(document, item)

    document.add_heading("Timing Definitions", level=2)
    add_table(
        document,
        ["Measurement", "Definition"],
        [
            ("CPU/OpenMP total", "Time from float input arrays entering the blend pipeline to the output float array being ready; exclude file I/O."),
            ("Blur", "All Gaussian blur work across source images and mask for all pyramid levels."),
            ("Downsample", "Pixel decimation after blur, excluding blur arithmetic."),
            ("Expand", "Upsampling used while constructing Laplacian bands."),
            ("Subtract", "Gaussian level minus expanded coarser level."),
            ("Blend", "Mask-weighted combination of corresponding Laplacian bands."),
            ("Collapse", "Expansion and addition while reconstructing the output."),
            ("CUDA transfers", "Host-to-device and device-to-host copy times measured separately with CUDA events."),
        ],
        widths=[1.45, 5.65],
    )

    document.add_heading("Metrics and Derived Results", level=1)
    for item in [
        "OpenMP speedup = serial total time divided by OpenMP total time.",
        "CUDA transfer-inclusive time = host-to-device + CUDA kernel + device-to-host time.",
        "CUDA kernel-only speedup = serial total time divided by CUDA kernel time.",
        "CUDA transfer-inclusive speedup = serial total time divided by transfer-inclusive CUDA time.",
        "Throughput = image width × image height divided by runtime, reported in megapixels per second.",
        "Report repeated-run average or median together with variation such as standard deviation.",
    ]:
        add_bullet(document, item)

    document.add_heading("Correctness Evaluation", level=1)
    for item in [
        "Collapse a CPU Laplacian pyramid and report maximum absolute reconstruction error.",
        "Compare OpenMP and CUDA outputs against the serial CPU reference.",
        "Report maximum absolute error, mean squared error, and optionally PSNR.",
        "Include visual output comparisons for representative masks and image pairs.",
        "Document that no objective blended-image ground truth is assumed; the serial CPU result is the numerical reference.",
    ]:
        add_bullet(document, item)

    document.add_heading("Recommended Figures and Tables", level=1)
    add_table(
        document,
        ["Type", "Recommended content"],
        [
            ("Figure", "Laplacian blending pipeline from inputs through pyramids, per-level blend, and collapse."),
            ("Figure", "Example Gaussian and Laplacian pyramid levels."),
            ("Figure", "Serial versus OpenMP runtime and speedup."),
            ("Figure", "CUDA transfer time versus kernel time."),
            ("Figure", "Per-stage stacked runtime chart."),
            ("Figure", "Runtime versus Gaussian kernel size."),
            ("Table", "Hardware and software configuration."),
            ("Table", "Numerical agreement between backends."),
            ("Table", "Runtime, speedup, and throughput results."),
        ],
        widths=[1.1, 6.0],
    )

    document.add_heading("Limitations and Reproducibility", level=1)
    for item in [
        "CUDA performance depends on GPU architecture, driver, CUDA version, and available memory.",
        "Preallocating all pyramid buffers increases memory usage but avoids repeated device allocation during the pipeline.",
        "Automatic resizing uses image A as the spatial reference and may distort image B when aspect ratios differ.",
        "Kernel fusion can improve performance but may reduce the precision of per-stage profiling; describe fused stages explicitly.",
        "Record compiler flags, OpenMP thread settings, GPU model, image dimensions, channel count, kernel size, levels, and repetitions.",
    ]:
        add_bullet(document, item)

    document.add_heading("IEEE Formatting Checklist", level=1)
    for item in [
        "Use the official IEEE conference two-column template for the final manuscript.",
        "Use numbered sections and subsections with concise descriptive headings.",
        "Place figure captions below figures and table captions above tables.",
        "Use numbered IEEE citations such as [1] and [2].",
        "Keep equations, units, abbreviations, and terminology consistent throughout.",
        "Do not report speedups or conclusions without the corresponding raw timing data and experimental conditions.",
    ]:
        add_bullet(document, item)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Laplacian Pyramid Blending Comparative Study")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
